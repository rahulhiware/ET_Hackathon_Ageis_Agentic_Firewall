import io
import json
import mimetypes
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Dict, Any

import requests
from bs4 import BeautifulSoup, Comment
import pymupdf
from docx import Document
from PIL import Image
import pytesseract

# Auto-detect Tesseract OCR binary on Windows if installed in common locations
for _cand in [
    Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
    Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    Path.home() / r"AppData\Local\Programs\Tesseract-OCR\tesseract.exe",
]:
    if _cand.exists():
        pytesseract.pytesseract.tesseract_cmd = str(_cand)
        break


def _html_to_text(raw_html: str) -> str:
    soup = BeautifulSoup(raw_html, "html.parser")
    parts = []

    visible = soup.get_text(" ", strip=True)
    if visible:
        parts.append("[VISIBLE_HTML]\n" + visible)

    comments = [
        str(c).strip()
        for c in soup.find_all(string=lambda x: isinstance(x, Comment))
        if str(c).strip()
    ]
    if comments:
        parts.append("[HTML_COMMENTS]\n" + "\n".join(comments))

    hidden = []
    for element in soup.find_all(True):
        style = (element.get("style") or "").lower().replace(" ", "")
        aria_hidden = (element.get("aria-hidden") or "").lower()
        if "display:none" in style or "visibility:hidden" in style or aria_hidden == "true":
            txt = element.get_text(" ", strip=True)
            if txt:
                hidden.append(txt)
    if hidden:
        parts.append("[HIDDEN_HTML]\n" + "\n".join(hidden))

    return "\n\n".join(parts)


def load_url(url: str) -> Dict[str, Any]:
    response = requests.get(
        url,
        timeout=15,
        headers={"User-Agent": "Aegis-Prompt-Injection-Firewall/1.0"},
    )
    response.raise_for_status()
    content_type = response.headers.get("content-type", "").lower()
    text = response.text
    if "html" in content_type or "<html" in text.lower():
        extracted = _html_to_text(text)
    else:
        extracted = text
    return {"source_type": "web", "name": url, "text": extracted}


def load_bytes(filename: str, data: bytes) -> Dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    mime = mimetypes.guess_type(filename)[0] or ""

    if suffix == ".pdf":
        doc = pymupdf.open(stream=data, filetype="pdf")
        pages = []
        for i, page in enumerate(doc):
            pages.append(f"[PDF PAGE {i+1}]\n{page.get_text()}")
        return {"source_type": "pdf", "name": filename, "text": "\n\n".join(pages)}

    if suffix == ".docx":
        doc = Document(io.BytesIO(data))
        chunks = [f"[PARAGRAPH]\n{p.text}" for p in doc.paragraphs if p.text.strip()]
        for ti, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text for cell in row.cells))
            chunks.append(f"[TABLE {ti+1}]\n" + "\n".join(rows))
        return {"source_type": "docx", "name": filename, "text": "\n\n".join(chunks)}

    if suffix in {".html", ".htm"} or "html" in mime:
        return {"source_type": "html", "name": filename, "text": _html_to_text(data.decode("utf-8", errors="replace"))}

    if suffix == ".eml":
        msg = BytesParser(policy=policy.default).parsebytes(data)
        parts = [f"[SUBJECT]\n{msg.get('subject', '')}", f"[FROM]\n{msg.get('from', '')}"]
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() in {"text/plain", "text/html"}:
                    try:
                        body = part.get_content()
                    except Exception:
                        continue
                    if part.get_content_type() == "text/html":
                        body = _html_to_text(body)
                    parts.append(body)
        else:
            try:
                parts.append(msg.get_content())
            except Exception:
                pass
        return {"source_type": "email", "name": filename, "text": "\n\n".join(parts)}

    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".gif"} or mime.startswith("image/"):
        image = Image.open(io.BytesIO(data))
        try:
            text = pytesseract.image_to_string(image)
        except pytesseract.TesseractNotFoundError:
            raise RuntimeError(
                "Tesseract OCR engine is not installed or not in PATH. "
                "To enable image OCR on Windows, install Tesseract OCR using: "
                "`winget install UB-Mannheim.TesseractOCR` and restart your terminal."
            )
        return {"source_type": "ocr", "name": filename, "text": text}

    # Markdown, source code, JSON/API responses, TXT and unknown textual files.
    text = data.decode("utf-8", errors="replace")
    if suffix in {".md", ".markdown"}:
        source_type = "markdown"
    elif suffix in {".py", ".js", ".ts", ".java", ".sql", ".sh", ".yaml", ".yml", ".json", ".xml"}:
        source_type = "source_code" if suffix not in {".json", ".xml"} else "api"
    else:
        source_type = "text"

    return {"source_type": source_type, "name": filename, "text": text}


def load_text(text: str, source_type: str) -> Dict[str, Any]:
    return {"source_type": source_type, "name": "inline", "text": text}
